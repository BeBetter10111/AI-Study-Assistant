import logging

import docx
import fitz
from pptx import Presentation


def extract_text(file_path: str, mime_type: str) -> str:
    extracted_text = []
    
    try:
        # Xử lý file PDF
        if "pdf" in mime_type.lower() or file_path.lower().endswith('.pdf'):
            with fitz.open(file_path) as doc:
                for page in doc:
                    # Trích xuất text cơ bản (Nếu cần OCR ảnh, có thể tích hợp pytesseract sau)
                    text = page.get_text("text")
                    if text:
                        extracted_text.append(text)
                        
        # Xử lý file Word (DOCX)
        elif "word" in mime_type.lower() or file_path.lower().endswith('.docx'):
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    extracted_text.append(para.text)
                    
        # Xử lý file PowerPoint (PPTX)
        elif "presentation" in mime_type.lower() or file_path.lower().endswith('.pptx'):
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        extracted_text.append(shape.text)
                        
        else:
            raise ValueError(f"Định dạng file không được hỗ trợ: {mime_type}")
            
        # Gộp tất cả các mảng thành một chuỗi văn bản hoàn chỉnh
        final_text = "\n".join(extracted_text)
        
        if not final_text.strip():
            logging.warning(f"Cảnh báo: Không tìm thấy nội dung văn bản trong {file_path}")
            
        return final_text.strip()
        
    except Exception as e:
        logging.error(f"Lỗi nghiêm trọng khi trích xuất text từ {file_path}: {str(e)}")
        raise e